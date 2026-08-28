# -*- coding: utf-8 -*-
"""Download the real datasets the project can legitimately use.

Everything the models train on was generated until now. This script fetches the
public sources that carry real observations, so the generated corpora can be
either replaced or corrected against them.

    python ml/fetch_real_data.py            (from the repo root)
    python ml/fetch_real_data.py --only nhamcs

Sources, all public and redistributable with attribution:

    nhamcs        US CDC National Hospital Ambulatory Medical Care Survey.
                  Real emergency department visits with a nurse-assigned
                  5-level triage immediacy, coded reasons for visit and
                  vital signs. No registration required.

    healthsites   OpenStreetMap health facilities for Bangladesh via HDX.
                  Real hospitals, pharmacies and labs with coordinates.

    dengue        Bangladesh dengue counts 2001-2024 compiled from DGHS
                  press releases, with the meteorology of each year.

    bangla_sx     Bangla disease-symptom associations from KUET: 172 Bangla
                  symptom names across 85 diseases.

Each download is cached under ``data/real/_cache`` so re-running is cheap and
works offline once primed. Nothing here contains patient-identifiable data:
NHAMCS is a de-identified national survey, the rest are reference data.
"""
from __future__ import annotations

import argparse
import csv
import io
import json
import re
import sys
import urllib.error
import urllib.request
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
REAL = ROOT / "data" / "real"
CACHE = REAL / "_cache"

# Some government hosts reject the default urllib agent.
USER_AGENT = "SasthoSetu/1.0 (+https://sasthosetu.gov.bd)"

NHAMCS_DATA = (
    "https://ftp.cdc.gov/pub/Health_Statistics/NCHS/Datasets/NHAMCS/ed2022.zip"
)
NHAMCS_DOC = (
    "https://ftp.cdc.gov/pub/Health_Statistics/NCHS/Dataset_Documentation/"
    "NHAMCS/doc22-ed-508.pdf"
)
HEALTHSITES = (
    "https://data.humdata.org/dataset/ab89f238-af45-419e-94aa-f91ef0ce42d0/"
    "resource/b639ee55-47f9-483c-a06e-ac2e8c12a1c0/download/bangladesh.csv"
)
DENGUE = (
    "https://raw.githubusercontent.com/awnonbhowmik/DENV-Data-Analysis/"
    "main/data/DENV_Bangladesh_2001_2024.xlsx"
)
BANGLA_SX = "https://data.mendeley.com/public-api/zip/rjgjh8hgrt/download/5"

# Fixed-width column positions in the NHAMCS ED file, 1-based inclusive as
# printed in the documentation. Converted to slices on read.
NHAMCS_FIELDS = {
    "age_years": (16, 18),
    "sex": (25, 25),
    "temperature_f": (48, 51),
    "pulse": (52, 54),
    "resp_rate": (55, 57),
    "bp_systolic": (58, 60),
    "bp_diastolic": (61, 63),
    "pulse_ox": (64, 66),
    "triage_level": (67, 68),
    "rfv1": (73, 77),
    "rfv2": (78, 82),
    "rfv3": (83, 87),
}

SEX_CODES = {"1": "F", "2": "M"}

# The survey codes immediacy 1 (seen immediately) to 5 (2-24 hours). Our own
# scale runs the other way round: 5 is the emergency. The mapping is a
# reversal, not a reinterpretation.
IMMEDIACY_TO_SEVERITY = {1: 5, 2: 4, 3: 3, 4: 2, 5: 1}


def _get(url: str, name: str) -> bytes:
    CACHE.mkdir(parents=True, exist_ok=True)
    cached = CACHE / name
    if cached.exists() and cached.stat().st_size > 0:
        print(f"  cached  {name} ({cached.stat().st_size:,} bytes)")
        return cached.read_bytes()

    print(f"  fetching {name} ...", flush=True)
    request = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with urllib.request.urlopen(request, timeout=300) as response:
        payload = response.read()
    cached.write_bytes(payload)
    print(f"  saved   {name} ({len(payload):,} bytes)")
    return payload


def _field(line: str, span: tuple[int, int]) -> str:
    start, end = span
    return line[start - 1:end].strip()


def _number(raw: str) -> float | None:
    """NHAMCS uses -9, -8, -7 for blank, unknown and not applicable."""
    if not raw:
        return None
    try:
        value = float(raw)
    except ValueError:
        return None
    return None if value < 0 else value


def _rfv_labels(doc_pdf: bytes) -> dict[str, str]:
    """Read the reason-for-visit code list out of the survey documentation.

    The codes in the data file are numeric; the documentation carries their
    plain-English wording, which is what makes the rows usable as text.
    """
    try:
        from pypdf import PdfReader
    except ImportError:
        print("  pypdf is not installed; reason codes will stay numeric")
        return {}

    reader = PdfReader(io.BytesIO(doc_pdf))
    text = "\n".join((page.extract_text() or "") for page in reader.pages)

    labels: dict[str, str] = {}
    for code, label in re.findall(r"^\s*(\d{4}\.\d)\s+([A-Z][^\n]{2,90})$",
                                  text, re.M):
        cleaned = " ".join(label.split()).rstrip(".")
        # The file stores 1050.1 as 10501, without the decimal point.
        labels.setdefault(code.replace(".", ""), cleaned)
    return labels


def fetch_nhamcs() -> Path:
    """Real ED visits with a clinician-assigned triage level."""
    print("\n=== NHAMCS: US emergency department visits")
    archive = _get(NHAMCS_DATA, "nhamcs_ed2022.zip")
    doc = _get(NHAMCS_DOC, "nhamcs_doc2022.pdf")

    labels = _rfv_labels(doc)
    print(f"  reason-for-visit labels: {len(labels)}")

    with zipfile.ZipFile(io.BytesIO(archive)) as zf:
        member = zf.namelist()[0]
        raw = zf.read(member).decode("latin-1")

    REAL.mkdir(parents=True, exist_ok=True)
    out = REAL / "nhamcs_ed_triage.csv"

    kept = 0
    skipped = 0
    with out.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([
            "age_years", "sex", "temperature_c", "pulse", "resp_rate",
            "bp_systolic", "bp_diastolic", "pulse_ox",
            "reason_1", "reason_2", "reason_3", "severity_level",
        ])

        for line in raw.splitlines():
            if not line.strip():
                continue

            immediacy = _number(_field(line, NHAMCS_FIELDS["triage_level"]))
            if immediacy is None or int(immediacy) not in IMMEDIACY_TO_SEVERITY:
                skipped += 1
                continue

            reasons = []
            for key in ("rfv1", "rfv2", "rfv3"):
                code = _field(line, NHAMCS_FIELDS[key])
                if code and not code.startswith("-"):
                    reasons.append(labels.get(code, ""))
            reasons = [r for r in reasons if r]
            if not reasons:
                skipped += 1
                continue

            fahrenheit = _number(_field(line, NHAMCS_FIELDS["temperature_f"]))
            # Recorded in tenths of a degree, so 986 is 98.6 F.
            celsius = (
                round(((fahrenheit / 10) - 32) * 5 / 9, 1)
                if fahrenheit and fahrenheit > 500 else None
            )

            writer.writerow([
                _number(_field(line, NHAMCS_FIELDS["age_years"])) or "",
                SEX_CODES.get(_field(line, NHAMCS_FIELDS["sex"]), ""),
                celsius if celsius is not None else "",
                _number(_field(line, NHAMCS_FIELDS["pulse"])) or "",
                _number(_field(line, NHAMCS_FIELDS["resp_rate"])) or "",
                _number(_field(line, NHAMCS_FIELDS["bp_systolic"])) or "",
                _number(_field(line, NHAMCS_FIELDS["bp_diastolic"])) or "",
                _number(_field(line, NHAMCS_FIELDS["pulse_ox"])) or "",
                reasons[0],
                reasons[1] if len(reasons) > 1 else "",
                reasons[2] if len(reasons) > 2 else "",
                IMMEDIACY_TO_SEVERITY[int(immediacy)],
            ])
            kept += 1

    print(f"  wrote {out.relative_to(ROOT)}: {kept:,} visits "
          f"({skipped:,} without a usable triage level or reason)")
    return out


def fetch_healthsites() -> Path:
    """Real Bangladeshi facilities with coordinates."""
    print("\n=== Healthsites: Bangladesh facilities (OpenStreetMap via HDX)")
    payload = _get(HEALTHSITES, "bd_healthsites.csv")
    rows = list(csv.DictReader(io.StringIO(payload.decode("utf-8"))))

    REAL.mkdir(parents=True, exist_ok=True)
    out = REAL / "bd_health_facilities.csv"

    counts: dict[str, int] = {}
    with out.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow([
            "name", "kind", "latitude", "longitude",
            "operator", "beds", "emergency", "speciality",
        ])
        for row in rows:
            name = (row.get("name") or "").strip()
            lat = (row.get("Y") or "").strip()
            lon = (row.get("X") or "").strip()
            if not name or not lat or not lon:
                continue

            kind = (row.get("healthcare") or row.get("amenity") or "").strip()
            if kind not in {"hospital", "clinic", "pharmacy", "laboratory",
                            "doctor", "doctors", "dentist"}:
                continue
            if kind in {"doctor", "doctors"}:
                kind = "doctor"

            writer.writerow([
                name, kind, lat, lon,
                (row.get("operator") or "").strip(),
                (row.get("beds") or "").strip(),
                (row.get("emergency") or "").strip(),
                (row.get("speciality") or "").strip(),
            ])
            counts[kind] = counts.get(kind, 0) + 1

    print(f"  wrote {out.relative_to(ROOT)}: "
          + ", ".join(f"{v:,} {k}" for k, v in sorted(counts.items())))
    return out


def fetch_dengue() -> Path:
    """Real dengue counts, replacing the generated surveillance series."""
    print("\n=== Dengue: Bangladesh 2001-2024 (compiled from DGHS releases)")
    payload = _get(DENGUE, "bd_dengue.xlsx")

    try:
        import openpyxl
    except ImportError:
        print("  openpyxl is not installed; skipping")
        return REAL / "bd_dengue_monthly.csv"

    book = openpyxl.load_workbook(io.BytesIO(payload), data_only=True)
    REAL.mkdir(parents=True, exist_ok=True)

    out = REAL / "bd_dengue_monthly.csv"
    sheet = book["Monthly_Infections_Long"]
    written = 0
    with out.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["year_month", "cases"])
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None:
                continue
            writer.writerow([row[0], row[1] if row[1] is not None else 0])
            written += 1
    print(f"  wrote {out.relative_to(ROOT)}: {written} months")

    regional = REAL / "bd_dengue_by_division.csv"
    sheet = book["Regional"]
    written = 0
    with regional.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["year", "division", "cases", "deaths"])
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if row[0] is None or row[1] is None:
                continue
            writer.writerow([row[0], row[1], row[2] or 0, row[3] or 0])
            written += 1
    print(f"  wrote {regional.relative_to(ROOT)}: {written} rows")
    return out


def fetch_bangla_symptoms() -> Path:
    """Bangla symptom vocabulary, to widen the lexicon."""
    print("\n=== Bangla disease-symptom associations (KUET, CC BY 4.0)")
    payload = _get(BANGLA_SX, "bangla_symptoms.zip")

    try:
        import openpyxl
    except ImportError:
        print("  openpyxl is not installed; skipping")
        return REAL / "bangla_disease_symptoms.csv"

    with zipfile.ZipFile(io.BytesIO(payload)) as zf:
        member = next(n for n in zf.namelist() if n.endswith("dataset.xlsx"))
        book = openpyxl.load_workbook(io.BytesIO(zf.read(member)))

    sheet = book.active
    symptoms = [c.value for c in sheet[1][1:]]

    REAL.mkdir(parents=True, exist_ok=True)
    out = REAL / "bangla_disease_symptoms.csv"
    diseases: set[str] = set()

    with out.open("w", newline="", encoding="utf-8") as fh:
        writer = csv.writer(fh)
        writer.writerow(["disease_bn", "symptoms_bn"])
        for row in range(2, sheet.max_row + 1):
            disease = sheet.cell(row, 1).value
            if not disease:
                continue
            present = [
                symptoms[i]
                for i in range(len(symptoms))
                if sheet.cell(row, i + 2).value in (1, 1.0, "1")
            ]
            if not present:
                continue
            writer.writerow([disease, "|".join(str(s) for s in present)])
            diseases.add(disease)

    print(f"  wrote {out.relative_to(ROOT)}: "
          f"{len(diseases)} diseases, {len(symptoms)} symptom columns")
    return out


SOURCES = {
    "nhamcs": fetch_nhamcs,
    "healthsites": fetch_healthsites,
    "dengue": fetch_dengue,
    "bangla_sx": fetch_bangla_symptoms,
}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--only",
        choices=sorted(SOURCES),
        action="append",
        help="fetch just this source (repeatable)",
    )
    args = parser.parse_args()

    wanted = args.only or sorted(SOURCES)
    failures = []

    for name in wanted:
        try:
            SOURCES[name]()
        except (urllib.error.URLError, TimeoutError, OSError) as error:
            # A source being unreachable must not stop the others: the caller
            # may be offline or behind a filter, and the cache may already
            # hold what they need.
            print(f"  FAILED {name}: {error}")
            failures.append(name)

    manifest = REAL / "MANIFEST.json"
    entries = {}
    for path in sorted(REAL.glob("*.csv")):
        with path.open(encoding="utf-8") as fh:
            rows = sum(1 for _ in fh) - 1
        entries[path.name] = {"rows": rows, "bytes": path.stat().st_size}
    manifest.write_text(json.dumps(entries, indent=2) + "\n", encoding="utf-8")
    print(f"\nmanifest -> {manifest.relative_to(ROOT)}")

    if failures:
        print(f"\n{len(failures)} source(s) unavailable: {', '.join(failures)}")
        sys.exit(1)


if __name__ == "__main__":
    main()
